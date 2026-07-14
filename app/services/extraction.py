"""
Text extraction service — converts uploaded files to plain text strings.

Supports: PDF (.pdf), Word (.docx), plain text (.txt), CSV (.csv), Excel (.xlsx).
Returns an empty string (never raises) if a file has no extractable text,
so the ingestion pipeline can emit a warning rather than crash (FR-2.6).
"""
import csv
import io
import logging
from pathlib import Path

logger = logging.getLogger(__name__)


def _extract_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
        reader = PdfReader(str(path))
        parts = []
        for page in reader.pages:
            text = page.extract_text()
            if text:
                parts.append(text.strip())
        return "\n\n".join(parts)
    except Exception as exc:
        logger.warning("PDF extraction failed for %s: %s", path.name, exc)
        return ""


def _extract_docx(path: Path) -> str:
    try:
        from docx import Document as DocxDocument
        doc = DocxDocument(str(path))
        parts = [para.text for para in doc.paragraphs if para.text.strip()]
        return "\n\n".join(parts)
    except Exception as exc:
        logger.warning("DOCX extraction failed for %s: %s", path.name, exc)
        return ""


def _extract_txt(path: Path) -> str:
    try:
        return path.read_text(encoding="utf-8", errors="replace").strip()
    except Exception as exc:
        logger.warning("TXT extraction failed for %s: %s", path.name, exc)
        return ""


def _extract_csv(path: Path) -> str:
    try:
        rows = []
        with path.open(encoding="utf-8", errors="replace", newline="") as f:
            reader = csv.reader(f)
            for row in reader:
                rows.append(", ".join(row))
        return "\n".join(rows)
    except Exception as exc:
        logger.warning("CSV extraction failed for %s: %s", path.name, exc)
        return ""


def _extract_xlsx(path: Path) -> str:
    try:
        import openpyxl
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        parts = []
        for sheet in wb.worksheets:
            rows = []
            for row in sheet.iter_rows(values_only=True):
                cells = [str(c) if c is not None else "" for c in row]
                if any(c.strip() for c in cells):
                    rows.append(", ".join(cells))
            if rows:
                parts.append(f"[Sheet: {sheet.title}]\n" + "\n".join(rows))
        wb.close()
        return "\n\n".join(parts)
    except Exception as exc:
        logger.warning("XLSX extraction failed for %s: %s", path.name, exc)
        return ""


# Public dispatcher
_EXTRACTORS = {
    ".pdf": _extract_pdf,
    ".docx": _extract_docx,
    ".doc": _extract_docx,
    ".txt": _extract_txt,
    ".csv": _extract_csv,
    ".xlsx": _extract_xlsx,
}


def extract_text(stored_path: str) -> str:
    """
    Extract plain text from a stored document file.

    Returns an empty string if the file cannot be read or has no text.
    Never raises — callers should check for empty string and emit a warning.
    """
    path = Path(stored_path)
    suffix = path.suffix.lower()
    extractor = _EXTRACTORS.get(suffix)
    if extractor is None:
        logger.warning("No extractor registered for extension %s", suffix)
        return ""
    return extractor(path)
